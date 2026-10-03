# -*- coding: utf-8 -*-
"""导出 30 条原文声称值抽取样本，供人工逐句核对（分层抽样）。"""
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE = 'audit_out_clean'
d = pd.read_csv(f'{BASE}/原文声称值_抽取明细.csv')
p = pd.read_csv(f'{BASE}/原文声称值vs重算值.csv').set_index('sig_id')

FIELDS = {'claim_c': ('C-index', None), 'claim_c_train': ('C-index', None), 'claim_c_valid': ('C-index', None),
          'claim_auc5': ('AUC', 5.0), 'claim_auc3': ('AUC', 3.0), 'claim_auc1': ('AUC', 1.0),
          'claim_auc_any': ('AUC', None), 'claim_auc5_max': ('AUC', 5.0), 'claim_auc5_valid': ('AUC', 5.0)}

d['used_field'] = ''
for i, r in d.iterrows():
    if r.sig_id not in p.index: continue
    pr = p.loc[r.sig_id]
    for f, (m, y) in FIELDS.items():
        v = pr.get(f)
        if pd.isna(v): continue
        if r.metric == m and abs(r.value - v) < 0.006 and (y is None or (not pd.isna(r.year) and abs(r.year - y) < 0.1)):
            d.at[i, 'used_field'] = f; break
d['used'] = (d.used_field != '').astype(int)
# 去重：同一签名+同一句+同一数值只保留一条，避免抽查位被重复句占满
dd = d.drop_duplicates(subset=['sig_id', 'sent', 'value']).copy()

# ---- 分层 ----
rng = np.random.default_rng(42)
L1 = dd[(dd.used_field == 'claim_auc5') & dd.sig_id.isin(p.index[p['偏倚_AUC5'].notna()])]
L2 = dd[dd.used_field.isin(['claim_c', 'claim_c_train', 'claim_c_valid']) & dd.sig_id.isin(p.index[p['偏倚_C'].notna()])]
risky = dd.bind.str.contains('one_nearest', na=False) | (dd.conf == 'win')
L3 = dd[risky]
L4 = dd[(dd.used == 1) & dd.sent.str.contains(r'1-\s*,?\s*3-', regex=True, na=False)]

def take(df, n):
    if len(df) == 0: return df
    idx = rng.choice(df.index.values, size=min(n, len(df)), replace=False)
    return df.loc[sorted(idx)]

parts = [('① 主结论支柱(5年AUC配对)', take(L1, 12)),
         ('② C-index配对', take(L2, 6)),
         ('③ 高风险绑定/窗口级', take(L3, 8)),
         ('④ 多年份连排(1-,3-,5-)', take(L4, 4))]
seen, rows = set(), []
for lab, sub in parts:
    for _, r in sub.iterrows():
        if r.name in seen: continue
        seen.add(r.name); rows.append((lab, r))
    # 补齐
need = {'① 主结论支柱(5年AUC配对)': 12, '② C-index配对': 6, '③ 高风险绑定/窗口级': 8, '④ 多年份连排(1-,3-,5-)': 4}[lab]
while sum(1 for x in rows if x[0] == lab) < need:
    pool = dd[~dd.index.isin(seen) & (dd.used == 1)]
    if len(pool) == 0: break
    i = rng.choice(pool.index.values)
    seen.add(i); rows.append((lab, dd.loc[i]))

REC = {'claim_auc5': ('重算AUC5中位', '偏倚_AUC5'), 'claim_auc5_max': ('重算AUC5中位', '偏倚_AUC5'),
       'claim_auc5_valid': ('重算AUC5中位', '偏倚_AUC5'), 'claim_c': ('重算C中位', '偏倚_C'),
       'claim_c_train': ('重算C中位', '偏倚_C'), 'claim_c_valid': ('重算C中位', '偏倚_C')}

wb = Workbook(); ws = wb.active; ws.title = '抽查30条'
head = ['序号', '抽样层', '签名ID', 'PMCID', '原文链接(点击打开)', '指标', '被采用字段\n(影响哪条结论)', '年份',
        '原文声称值', '我方重算值', '偏倚(原文−重算)', '抽取方式', '置信度', '语境', '原文原句（请核对这句里这个值）',
        '你的核对结果\n(对/错/存疑)', '备注']
ws.append(head)
hf = PatternFill('solid', fgColor='D9E2F3'); thin = Side(style='thin', color='BFBFBF')
for c in ws[1]:
    c.font = Font(bold=True, size=10); c.fill = hf
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = Border(bottom=thin)
url_font = Font(color='0563C1', underline='single', size=10)
for n, (lab, r) in enumerate(rows, 1):
    pr = p.loc[r.sig_id] if r.sig_id in p.index else None
    if r.used_field in REC:
        rc, bc = REC[r.used_field]
        rv, bv = (pr[rc], pr[bc]) if pr is not None else (np.nan, np.nan)
    else:
        rv, bv = np.nan, np.nan
    url = f'https://europepmc.org/article/PMC/{r.pmcid}'
    ws.append([n, lab, r.sig_id, r.pmcid, url, r.metric, r.used_field or '（未采用）',
               '' if pd.isna(r.year) else int(r.year), r.value,
               '' if pd.isna(rv) else round(float(rv), 4),
               '' if pd.isna(bv) else round(float(bv), 4),
               r.bind, r.conf, r.ctx, r.sent, '', ''])
    row = ws.max_row
    ws.cell(row, 5).hyperlink = url; ws.cell(row, 5).font = url_font
    ws.cell(row, 15).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row, 16).fill = PatternFill('solid', fgColor='FFF2CC')
    ws.cell(row, 10).number_format = '0.000'; ws.cell(row, 11).number_format = '0.000'
    ws.row_dimensions[row].height = 58
W = [5, 24, 9, 13, 34, 8, 16, 6, 10, 11, 12, 13, 9, 7, 85, 14, 16]
for i, w in enumerate(W, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:Q{ws.max_row}'

# 说明页
ws2 = wb.create_sheet('核对说明')
notes = [
 ['核对说明', ''],
 ['目的', '验证脚本从 319 篇全文自动抽取的「原文声称 C-index / AUC」是否抽对。这是全流程唯一未逐条人工核对的环节，估测准确率约 90%。'],
 ['抽样方式', '固定随机种子 42，分 4 层：① 决定主结论的 5 年 AUC 配对 12 条；② C-index 配对 6 条；③ 高风险抽取（one_nearest 绑定 / 窗口级置信）8 条；④ 1-,3-,5-year 连排句 4 条。'],
 ['怎么核', '点「原文链接」打开 PMC 全文 → 在页面内搜索该数值（如 0.826）→ 看它对应的指标（AUC 还是 C-index）、年份（1/3/5 年）、队列（train 还是 validation）是否与我方记录一致。'],
 ['判定', '在「你的核对结果」列填：对 / 错 / 存疑。错例请说明正确值。'],
 ['重点看什么', '1) 年份绑定：连排句 "1-, 3-, and 5-year ... 0.826, 0.764, 0.755" 必须一一对应，不能串位；2) 指标混淆：AUC 与 C-index 不能互串；3) 队列归属：train（训练）还是 validation（验证），我方只用 validation 才算严格外部验证；4) 数值所在句是否是该签名的最终模型，而非对比模型/其他模型。'],
 ['容错口径', '若 30 条里错误 ≤3 条（错误率 ≤10%），主结论（原文 5 年 AUC 比重算高 +0.174）稳健；若错误集中在「年份串位」，需回修 18_claim_extract.py 后重跑 Part 2 发表偏倚。'],
 ['字段释义', 'bind=抽取方式(seq_post 序列后绑定 / seq_pre 前绑定 / one_nearest 就近取一个，风险最高)；conf=置信度(sent 句子级，好 / win 窗口级，次好)；ctx=语境(train 训练队列 / valid 验证队列 / unk 未标明)。'],
]
for r in notes:
    ws2.append(r)
ws2.column_dimensions['A'].width = 14; ws2.column_dimensions['B'].width = 110
for row in ws2.iter_rows(min_row=2):
    row[0].font = Font(bold=True); row[1].alignment = Alignment(wrap_text=True, vertical='top')
    ws2.row_dimensions[row[0].row].height = 46
ws2['A1'].font = Font(bold=True, size=13)

out = f'{BASE}/人工抽查_30条原文声称值.xlsx'
wb.save(out)

csv_rows = []
for n, (lab, r) in enumerate(rows, 1):
    pr = p.loc[r.sig_id] if r.sig_id in p.index else None
    rv = bv = np.nan
    if r.used_field in REC:
        rc_, bc_ = REC[r.used_field]
        if pr is not None: rv, bv = pr[rc_], pr[bc_]
    csv_rows.append(dict(序号=n, 抽样层=lab, 签名ID=r.sig_id, PMCID=r.pmcid,
                         链接=f'https://europepmc.org/article/PMC/{r.pmcid}', 指标=r.metric,
                         被采用字段=r.used_field or '（未采用）', 年份='' if pd.isna(r.year) else int(r.year),
                         原文声称值=r.value, 我方重算值='' if pd.isna(rv) else round(float(rv), 4),
                         偏倚='' if pd.isna(bv) else round(float(bv), 4), 抽取方式=r.bind, 置信度=r.conf,
                         语境=r.ctx, 原文原句=r.sent, 你的核对结果='', 备注=''))
pd.DataFrame(csv_rows).to_csv(f'{BASE}/人工抽查_30条原文声称值.csv', index=False, encoding='utf-8-sig')
print('导出完成:', out)
print('分层计数:', pd.Series([x[0] for x in rows]).value_counts().to_string())
print('被采用记录占比:', sum(1 for _, r in rows if r.used == 1), '/', len(rows))
